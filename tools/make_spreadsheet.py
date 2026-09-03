#!/usr/bin/env python3
"""Rebuild data/Wondrous_Magic_script.xlsx from the ROM and the translation
modules, so the workbook is never out of step with what is actually inserted.

    python3 tools/make_spreadsheet.py
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import paths
import wmtool
import script_a, script_b, script_c
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

HEAD = PatternFill('solid', fgColor='1F3864')
EDIT = PatternFill('solid', fgColor='FFF9D6')
DONE = PatternFill('solid', fgColor='E7F3E7')
THIN = Side(style='thin', color='D0D0D0')

GLOSSARY = [
    ('リンクル', 'Rinkle', 'character', "The heroine's grandfather, once court mage of Aria."),
    ('シリア', 'Silia', 'character', "Her mother. Died of an unnamed sickness soon after the birth."),
    ('グナス', 'Gunas', 'character', "Her father. Left for Darles the day after Silia died."),
    ('ソルディック', 'Soldic', 'character', 'Court swordsman of Aria.'),
    ('シーラ', 'Sheila', 'character', 'Teacher at the magic guild; sets the Rose Coat examination.'),
    ('アルダン王', 'King Aldan', 'character', 'King of Aria.'),
    ('チャル', 'Charl', 'character', "The heroine's pet. Says \u201cKyui\u201d."),
    ('サダ', 'Sada', 'character', 'Villager; gives you wine when you leave.'),
    ('クルエル', 'Cruel', 'character', 'Villager.'),
    ('セグ', 'Seg', 'character', 'Villager, called Uncle Seg.'),
    ('ネル', 'Nel', 'character', 'Villager.'),
    ('リナ', 'Rina', 'character', "The heroine's friend."),
    ('フレディア', 'Fredia', 'character', ''),
    ('ガストール', 'Gastor', 'character', ''),
    ('ガスタック', 'Gastack', 'character', ''),
    ('ピリオド', 'Period', 'character', ''),
    ('セシリア', 'Cecilia', 'character', ''),
    ('ザイフォン', 'Zaifon', 'character', ''),
    ('イリオン', 'Ilion', 'character', ''),
    ('シュレル', 'Shrell', 'deity', 'The goddess. Gave magic to the chosen.'),
    ('イヴァス', 'Ivuas', 'deity', 'God of the phantom beasts. Cursed the 99th daughter.'),
    ('アリアの都', 'the city of Aria', 'place', 'Lower case "city"; the name alone is just "Aria".'),
    ('ダーレス', 'Darles', 'place', 'Where Gunas went.'),
    ('クローグ', 'Kroag', 'place', 'Town north of Aria, with the tower of magic.'),
    ('ネクストリア', 'Nextria', 'place', ''),
    ('魔法の塔', 'tower of magic', 'place', ''),
    ('魔法院', 'magic guild', 'place', 'Not "academy" or "institute".'),
    ('幻獣', 'phantom beast', 'term', 'Never "illusion beast" or "monster".'),
    ('魔法使い', 'mage', 'term', 'Not "magic user" or "wizard" - Wizard is a class name.'),
    ('女魔法使い', 'sorceress', 'term', 'Used for female-only equipment.'),
    ('戦士', 'warrior', 'term', ''),
    ('魔法戦士', 'magic warrior', 'term', ''),
    ('城の戦士', 'palace warrior', 'term', ''),
    ('巫女', 'priestess', 'term', ''),
    ('呪文', 'spell', 'term', 'The offensive half of magic.'),
    ('祈願 / 祈願力', 'prayer', 'term', 'The healing half, and its resource.'),
    ('体力', 'health', 'term', ''),
    ('経験値', 'EXP', 'term', ''),
    ('状態', 'condition', 'term', 'Status ailments; "Cures one rank of status".'),
    ('戦闘不能', 'fallen', 'term', 'KO in the status window, for width.'),
    ('鑑定', 'appraise', 'term', 'Unidentified items must be appraised.'),
    ('装備', 'equip', 'term', ''),
    ('防具', 'armour', 'term', 'British spelling throughout - armour, not armor.'),
    ('ベゼッタ', 'Bezetta', 'currency', ''),
    ('オニキス', 'Onyx', 'currency', ''),
    ('ローズコーツ', 'Rose Coat', 'item', 'The examination jewel in the tower at Kroag.'),
    ('ブラッドストーン', 'Bloodstone', 'item', 'Written in \u301d \u301e quotes in the script.'),
    ('シュレルの杖', "Shrell's Staff", 'item', 'Assembled from five pieces.'),
    ('ドラゴノイム', 'Dragonoim', 'creature', 'Called by the Sea Horn.'),
    ('ヴァンパイア', 'vampire', 'creature', ''),
    ('「', '"', 'punctuation', 'Opening quote; rendered as a plain double quote.'),
    ('…', '...', 'punctuation', 'Three periods. Keep them - they carry hesitation and grief.'),
    ('・', '.', 'punctuation', ''),
]


def english():
    out = dict(script_c.TEXT)
    text_a, missing = script_a.build()
    assert not missing, missing[:3]
    out.update(text_a)
    out.update(script_b.TEXT)
    return out


def main():
    en = english()
    wb = Workbook()

    ws = wb.active
    ws.title = 'Script'
    hdr = ['ID', 'Block', 'ROM', 'CPU', 'Bytes', 'Status', 'Japanese', 'English',
           'EN bytes', 'Raw hex']
    ws.append(hdr)
    rows = 0
    for blk, a, b in wmtool.TEXT_BLOCKS:
        p, n = a, 0
        while p < b:
            toks, q = wmtool.decode(p)
            bank, addr = wmtool.pc2snes(p)
            t = en.get(p, '')
            ws.append(['%s%04d' % (blk, n), blk, '$%06X' % p,
                       '$%02X:%04X' % (bank, addr), q - p,
                       'done' if t else '', wmtool.text(toks).rstrip('\n'), t,
                       None, wmtool.ROM[p:q].hex().upper()])
            rows += 1
            p, n = q, n + 1
    for r in range(2, rows + 2):
        ws.cell(r, 9).value = '=IF(H%d="","",LEN(H%d))' % (r, r)

    for c in range(1, len(hdr) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:J%d' % (rows + 1)
    for k, v in {'A': 9, 'B': 6, 'C': 10, 'D': 11, 'E': 7, 'F': 8,
                 'G': 46, 'H': 46, 'I': 9, 'J': 46}.items():
        ws.column_dimensions[k].width = v
    for r in range(2, rows + 2):
        for c in range(1, len(hdr) + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name='Arial', size=10)
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (7, 8)))
        ws.cell(r, 7).font = Font(name='Arial', size=11)
        ws.cell(r, 10).font = Font(name='Courier New', size=8, color='808080')
        ws.cell(r, 8).fill = DONE if ws.cell(r, 6).value else EDIT

    gs = wb.create_sheet('Glossary')
    gs.append(['Japanese', 'English', 'Type', 'Notes'])
    for row in GLOSSARY:
        gs.append(list(row))
    for c in range(1, 5):
        cell = gs.cell(1, c)
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal='center')
    for r in range(2, gs.max_row + 1):
        for c in range(1, 5):
            cell = gs.cell(r, c)
            cell.font = Font(name='Arial', size=11 if c == 1 else 10)
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical='top', wrap_text=(c == 4))
    gs.freeze_panes = 'A2'
    gs.auto_filter.ref = 'A1:D%d' % gs.max_row
    for col, w in zip('ABCD', (22, 22, 12, 62)):
        gs.column_dimensions[col].width = w

    cc = wb.create_sheet('Control codes')
    cc.append(['Byte', 'Token', 'Args', 'Meaning'])
    meta = {0x00: 'end of string', 0x01: 'no-op',
            0x02: "print the hero's name", 0x03: '7-digit decimal',
            0x04: '8-digit decimal', 0x05: 'wait for input',
            0x06: 'JSL $10:94F1', 0x07: 'JSL $10:94E7',
            0x08: 'scene continue', 0x09: 'JSL $10:94EB',
            0x0A: 'clear to a new page', 0x0B: 'sound effect',
            0x0C: 'JSL $00:FB7C', 0x0D: 'newline'}
    for b in sorted(wmtool.CTRL):
        cc.append(['$%02X' % b, wmtool.CTRL[b][0], wmtool.CTRL[b][1], meta[b]])
    cc.append(['$1E nn', 'dictionary', 1, 'expands dictionary entry nn'])
    cc.append(['$1E $FF nn', 'dictionary', 2, 'expands entry 256 + nn'])
    cc.append(['$7F', 'space', 0, 'half-width space'])
    cc.append(['$20', 'space', 0, 'full-width space, one whole cell'])

    lg = wb.create_sheet('Read me', 0)
    total = rows
    done = sum(1 for v in en)
    for line in [
        ['Wondrous Magic - translation worksheet'], [''],
        ['Generated by tools/make_spreadsheet.py. Do not hand-edit and expect it'],
        ['to survive - the English lives in build/script_a.py, script_b.py and'],
        ['script_c.py, and this workbook is rebuilt from those.'], [''],
        ['Strings', total], ['Translated', done],
        ['Remaining', total - done], [''],
        ['Sheet', 'Contents'],
        ['Script', 'Every string. Green English cells are inserted; yellow are open.'],
        ['Glossary', 'Agreed renderings. Check before inventing one.'],
        ['Control codes', 'Tokens that must survive translation, in order.'], [''],
        ['Line rules'],
        ['1', 'A dialogue box is 14 cells: 28 half-width characters per line.'],
        ['2', 'Four rows. Speaker on its own line, then three lines before <WAIT>.'],
        ['3', 'ASCII only. " for the Japanese quote, ... for the ellipsis.'],
        ['4', 'English must fit the Bytes column; the dictionary buys the headroom.'],
    ]:
        lg.append(line)
    lg.column_dimensions['A'].width = 14
    lg.column_dimensions['B'].width = 78
    lg['A1'].font = Font(name='Arial', size=14, bold=True)
    for row in lg.iter_rows():
        for cell in row:
            if not cell.font.bold:
                cell.font = Font(name='Arial', size=10)

    out = os.path.join(paths.DATA, 'Wondrous_Magic_script.xlsx')
    wb.save(out)
    print('%s: %d strings, %d translated, %d glossary entries'
          % (os.path.basename(out), total, done, len(GLOSSARY)))


if __name__ == '__main__':
    main()
